"""
Assembler tests (implementation.md P2 task 2.6). Covers the AI watch-list
items called out for this task explicitly: S06 (empty special clause) leaves
no leftover placeholder or double space; the .docx round-trips; RenderedDoc
and the written .docx agree (the Q8 property, checked properly in
test_qa_gate.py, but a basic version is worth having here too).
"""
from __future__ import annotations

import io
import re

import docx

from core.models import (
    BusinessInputs,
    ClauseAssessment,
    Confidence,
    FieldDecision,
    FieldStatus,
    Readiness,
    ReviewerAction,
    ReviewerActionType,
)
from tools import template_store
from tools.assembler import build_rendered_doc, write_docx


def _template():
    return template_store.get_template("mutual_nda")


def _inputs(**overrides):
    base = dict(
        disclosing_party="Zycus Inc.",
        receiving_party="Northwind Vendor Solutions Pvt. Ltd.",
        effective_date="use today's date",
        term="2 years from effective date",
        governing_law="State of Delaware, USA",
        survival_period="3 years after termination",
        purpose="Evaluating a potential vendor relationship for procurement software integration",
        special_clause="Receiving party wants a carve-out allowing disclosure to their affiliates without prior written consent",
        payment_terms="Not applicable — this is an NDA, not a commercial agreement",
        contract_value="",
    )
    base.update(overrides)
    return BusinessInputs(**base)


def _auto_filled(field, value, status=FieldStatus.AUTO_FILLED, gate="G11_default"):
    return FieldDecision(field=field, status=status, gate=gate, value_for_document=value, reason="ok")


def _sample_decisions():
    return [
        _auto_filled("disclosing_party", "Zycus Inc."),
        _auto_filled("receiving_party", "Northwind Vendor Solutions Pvt. Ltd."),
        _auto_filled("effective_date", "13 September 2026", status=FieldStatus.AUTO_FILLED_WITH_ASSUMPTION, gate="G10_assumption"),
        _auto_filled("term", "24 months"),
        _auto_filled("survival_period", "36 months"),
        _auto_filled("governing_law", "State of Delaware, USA"),
        _auto_filled("purpose", "Evaluating a potential vendor relationship for procurement software integration"),
        FieldDecision(field="special_clause", status=FieldStatus.NEEDS_REVIEW, gate="G5_risk_shifting",
                      value_for_document=None, reason="risk-shifting clause"),
        FieldDecision(field="payment_terms", status=FieldStatus.NOT_APPLICABLE, gate="G3_not_applicable",
                      value_for_document=None, reason="not applicable"),
        FieldDecision(field="contract_value", status=FieldStatus.NOT_APPLICABLE, gate="G3_not_applicable",
                      value_for_document=None, reason="not applicable"),
    ]


def _sample_clause():
    return ClauseAssessment(
        request_summary="Affiliate disclosure carve-out",
        library_match="affiliate_disclosure",
        deviation="non_standard",
        risks=["Confidential Information could reach non-signatory affiliates"],
        conflicting_sections=[5],
        rule_ids_cited=["NDA-DISC-01"],
        proposed_text="Notwithstanding the foregoing... (library text)",
        proposed_text_source="library",
        confidence=Confidence.HIGH,
    )


class TestS01Sample:
    def test_renders_without_error(self):
        rendered = build_rendered_doc(
            _template(), _inputs(), _sample_decisions(), _sample_clause(), Readiness.NEEDS_REVIEW,
        )
        assert rendered.title == "MUTUAL NON-DISCLOSURE AGREEMENT"
        assert rendered.banner is None

    def test_docx_round_trips(self):
        rendered = build_rendered_doc(
            _template(), _inputs(), _sample_decisions(), _sample_clause(), Readiness.NEEDS_REVIEW,
        )
        data = write_docx(rendered)
        doc = docx.Document(io.BytesIO(data))
        text = "\n".join(p.text for p in doc.paragraphs)
        assert "Zycus Inc." in text
        assert "Northwind Vendor Solutions Pvt. Ltd." in text
        assert "State of Delaware, USA" in text
        assert "library text" in text  # pending proposed clause shown by default

    def test_pending_clause_is_highlighted(self):
        rendered = build_rendered_doc(
            _template(), _inputs(), _sample_decisions(), _sample_clause(), Readiness.NEEDS_REVIEW,
        )
        section4 = rendered.paragraphs[4]  # preamble, 1, 2, 3, 4 -> index 4
        highlighted = [r for r in section4.runs if r.highlight == "yellow"]
        assert highlighted
        assert "library text" in highlighted[0].text


class TestS06EmptySpecialClause:
    def test_no_placeholder_or_double_space(self):
        decisions = _sample_decisions()
        for d in decisions:
            if d.field == "special_clause":
                d.status = FieldStatus.AUTO_FILLED
                d.gate = "G11_default"
                d.value_for_document = None
        rendered = build_rendered_doc(
            _template(), _inputs(special_clause=""), decisions, None, Readiness.READY_FOR_SIGNATURE_REVIEW,
        )
        full_text = rendered.full_text
        assert "[SPECIAL CLAUSE" not in full_text
        assert "  " not in full_text  # no double space anywhere
        assert not re.search(r"\S \. ", full_text)  # no dangling space before the period
        section4_text = rendered.paragraphs[4].text
        assert section4_text.rstrip().endswith("Section 5.")


class TestBlockedMissing:
    def test_missing_field_gets_marker_and_banner(self):
        decisions = _sample_decisions()
        for d in decisions:
            if d.field == "governing_law":
                d.status = FieldStatus.BLOCKED_MISSING
                d.gate = "G1_missing"
                d.value_for_document = None
        rendered = build_rendered_doc(
            _template(), _inputs(governing_law=""), decisions, _sample_clause(), Readiness.BLOCKED,
        )
        assert rendered.banner is not None
        assert "NOT READY" in rendered.banner
        assert "1" in rendered.banner
        assert "⟦MISSING: GOVERNING LAW⟧" in rendered.full_text
        assert "[GOVERNING LAW]" not in rendered.full_text


class TestReviewerActions:
    def _base(self):
        return _template(), _inputs(), _sample_decisions(), _sample_clause()

    def test_accept_proposed_removes_highlight(self):
        template, inputs, decisions, clause = self._base()
        action = ReviewerAction(field="special_clause", action=ReviewerActionType.ACCEPT_PROPOSED)
        rendered = build_rendered_doc(template, inputs, decisions, clause, Readiness.NEEDS_REVIEW, [action])
        section4 = rendered.paragraphs[4]
        assert any("library text" in r.text for r in section4.runs)
        assert all(r.highlight != "yellow" for r in section4.runs if "library text" in r.text)

    def test_remove_leaves_nothing(self):
        template, inputs, decisions, clause = self._base()
        action = ReviewerAction(field="special_clause", action=ReviewerActionType.REMOVE)
        rendered = build_rendered_doc(template, inputs, decisions, clause, Readiness.NEEDS_REVIEW, [action])
        assert "library text" not in rendered.full_text
        assert "affiliate" not in rendered.full_text.lower() or "Affiliate" not in rendered.full_text

    def test_use_as_requested_inserts_raw_text(self):
        template, inputs, decisions, clause = self._base()
        action = ReviewerAction(field="special_clause", action=ReviewerActionType.USE_AS_REQUESTED)
        rendered = build_rendered_doc(template, inputs, decisions, clause, Readiness.NEEDS_REVIEW, [action])
        assert inputs.special_clause in rendered.full_text

    def test_edit_inserts_edited_text(self):
        template, inputs, decisions, clause = self._base()
        action = ReviewerAction(field="special_clause", action=ReviewerActionType.EDIT,
                                 edited_text="Custom reviewer-edited carve-out language.")
        rendered = build_rendered_doc(template, inputs, decisions, clause, Readiness.NEEDS_REVIEW, [action])
        assert "Custom reviewer-edited carve-out language." in rendered.full_text


class TestPendingReviewNoProposal:
    def test_no_clause_and_no_reviewer_action_shows_pending_marker(self):
        decisions = _sample_decisions()
        rendered = build_rendered_doc(
            _template(), _inputs(), decisions, None, Readiness.NEEDS_REVIEW,
        )
        assert "⟦PENDING REVIEW: special clause⟧" in rendered.full_text
