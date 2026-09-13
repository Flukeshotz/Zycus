"""
Document Assembler (architecture.md §3 #14, §9, D-06, D-32, D-52).

Two stages, deliberately kept separate:

  1. build_rendered_doc(...) -> RenderedDoc
     A plain-Python intermediate representation (paragraphs -> runs, each
     run carrying a highlight/note flag). All the *decisions* about what
     goes where happen here, and this function has no python-docx
     dependency at all — easy to unit test, and reused as-is by the HTML
     preview (tools/preview.py, P4) so the preview and the .docx can never
     drift apart (D-52).

  2. write_docx(rendered) -> bytes
     A dumb, mechanical renderer: walks RenderedDoc and writes a .docx.
     No decisions happen here — if something needs to change about *what*
     appears in the document, that's a build_rendered_doc change, not a
     write_docx change.

The LLM never writes the document (D-06): only build_rendered_doc decides
final content, from the template's own text, the router's FieldDecisions,
and (for the one flagged clause) run_result.clause + a reviewer_action.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from typing import Literal

from core.models import (
    BusinessInputs,
    ClauseAssessment,
    FieldDecision,
    FieldStatus,
    Readiness,
    ReviewerAction,
    ReviewerActionType,
)
from tools.mapper import map_decisions_to_placeholders
from tools.rulebook import display_field_name
from tools.template_store import Template

Highlight = Literal["none", "yellow", "red"]

MISSING_MARKER_TEMPLATE = "⟦MISSING: {field}⟧"
PENDING_REVIEW_MARKER = "⟦PENDING REVIEW: special clause⟧"
SPECIAL_CLAUSE_PLACEHOLDER_NAME = "SPECIAL CLAUSE, IF ANY"
SPECIAL_CLAUSE_TOKEN = f"[{SPECIAL_CLAUSE_PLACEHOLDER_NAME}]"


# ---------------------------------------------------------------------------
# RenderedDoc — the shared intermediate model (D-52)
# ---------------------------------------------------------------------------


@dataclass
class RenderedRun:
    text: str
    highlight: Highlight = "none"
    note: str | None = None
    # "template" = verbatim text from the template file (or static
    # boilerplate like the signature block); "slot" = substituted content
    # (a field value, a MISSING/PENDING marker, or clause text). The QA
    # gate's Q2 and Q5 checks (architecture §9) are scoped by this tag —
    # forbidden literals are only checked in slots, and only "template" runs
    # are compared byte-for-byte against the source template.
    source: Literal["template", "slot"] = "template"


@dataclass
class RenderedParagraph:
    runs: list[RenderedRun] = field(default_factory=list)

    @property
    def text(self) -> str:
        return "".join(r.text for r in self.runs)


@dataclass
class RenderedDoc:
    title: str
    paragraphs: list[RenderedParagraph] = field(default_factory=list)
    banner: str | None = None  # "DRAFT — NOT READY: n required field(s) missing" (D-17)

    @property
    def full_text(self) -> str:
        # Order must match write_docx exactly (banner, then title, then body)
        # — this is what the QA gate's Q8 check compares against the
        # reopened .docx, so a mismatch here is a false QA failure, not a
        # real one (see docs/DEBUG_LOG.md).
        parts = []
        if self.banner:
            parts.append(self.banner)
        parts.append(self.title)
        parts.extend(p.text for p in self.paragraphs if p.text)
        return "\n\n".join(parts)


# ---------------------------------------------------------------------------
# build_rendered_doc
# ---------------------------------------------------------------------------

_PLACEHOLDER_TOKEN = re.compile(r"\[([^\[\]]+)\]")


def build_rendered_doc(
    template: Template,
    inputs: BusinessInputs,
    decisions: list[FieldDecision],
    clause: ClauseAssessment | None,
    readiness: Readiness,
    reviewer_actions: list[ReviewerAction] | None = None,
) -> RenderedDoc:
    decisions_by_field = {d.field: d for d in decisions}
    actions_by_field = {a.field: a for a in (reviewer_actions or [])}

    placeholder_runs = _simple_placeholder_runs(template, decisions_by_field)

    special_decision = decisions_by_field.get("special_clause")
    special_runs = _resolve_special_clause_runs(
        special_decision, clause, actions_by_field.get("special_clause"), inputs.special_clause,
    )
    placeholder_runs[SPECIAL_CLAUSE_PLACEHOLDER_NAME] = special_runs

    paragraphs: list[RenderedParagraph] = []
    for section in template.sections:
        section_text = template.section_text[section]
        if section == "4":
            runs = _build_section_4_runs(section_text, special_runs)
        else:
            runs = _substitute_section(section_text, placeholder_runs)
        paragraphs.append(RenderedParagraph(runs=runs))

    paragraphs.extend(_signature_block_paragraphs())

    banner = None
    if readiness == Readiness.BLOCKED:
        missing = sum(1 for d in decisions if d.status == FieldStatus.BLOCKED_MISSING)
        banner = f"DRAFT — NOT READY: {missing} required field{'s' if missing != 1 else ''} missing"

    return RenderedDoc(title=template.title, paragraphs=paragraphs, banner=banner)


def _simple_placeholder_runs(
    template: Template, decisions_by_field: dict[str, FieldDecision]
) -> dict[str, list[RenderedRun]]:
    """Build runs for every placeholder except SPECIAL CLAUSE, IF ANY, which
    has its own resolution logic (_resolve_special_clause_runs)."""
    result: dict[str, list[RenderedRun]] = {}
    for placeholder_name, spec in template.placeholders.items():
        if placeholder_name == SPECIAL_CLAUSE_PLACEHOLDER_NAME:
            continue
        decision = decisions_by_field.get(spec.field)
        result[placeholder_name] = _field_runs(spec.field, decision)
    return result


def _field_runs(field_name: str, decision: FieldDecision | None) -> list[RenderedRun]:
    if decision is None:
        return [RenderedRun(text="", source="slot")]
    if decision.status == FieldStatus.BLOCKED_MISSING:
        marker = MISSING_MARKER_TEMPLATE.format(field=display_field_name(field_name).upper())
        return [RenderedRun(text=marker, highlight="red", note=decision.reason, source="slot")]
    if decision.value_for_document:
        highlight: Highlight = "yellow" if decision.status == FieldStatus.NEEDS_REVIEW else "none"
        note = decision.reason if highlight == "yellow" else None
        return [RenderedRun(text=decision.value_for_document, highlight=highlight, note=note, source="slot")]
    return [RenderedRun(text="", source="slot")]


def _resolve_special_clause_runs(
    decision: FieldDecision | None,
    clause: ClauseAssessment | None,
    reviewer_action: ReviewerAction | None,
    raw_value: str,
) -> list[RenderedRun]:
    """
    The one field whose document content is never taken directly from
    FieldDecision.value_for_document (D-12). Never inserts unvetted request
    text without either a human decision or, by default, the safe proposed
    fallback (never the raw request itself while pending, D-10).
    """
    if decision is None or decision.status != FieldStatus.NEEDS_REVIEW:
        # AUTO_FILLED (empty clause, S06) or any other non-flagged state.
        return [RenderedRun(text="", source="slot")]

    if reviewer_action is not None:
        return _apply_reviewer_action(reviewer_action, clause, raw_value)

    # No human decision yet: show the safe default (proposed fallback),
    # highlighted and pending, or a marker if no proposal exists at all.
    if clause and clause.proposed_text:
        return [RenderedRun(
            text=clause.proposed_text, highlight="yellow",
            note=f"Pending review: {decision.reason}", source="slot",
        )]
    return [RenderedRun(
        text=PENDING_REVIEW_MARKER, highlight="yellow",
        note=f"Pending review: {decision.reason} AI analysis unavailable — draft this clause manually.",
        source="slot",
    )]


def _apply_reviewer_action(
    action: ReviewerAction, clause: ClauseAssessment | None, raw_value: str
) -> list[RenderedRun]:
    if action.action == ReviewerActionType.ACCEPT_PROPOSED:
        text = clause.proposed_text if clause and clause.proposed_text else ""
        if not text:
            return [RenderedRun(
                text=PENDING_REVIEW_MARKER, highlight="yellow",
                note="No proposed text was available to accept; draft this clause manually.", source="slot",
            )]
        return [RenderedRun(text=text, highlight="none", source="slot")]

    if action.action == ReviewerActionType.USE_AS_REQUESTED:
        return [RenderedRun(
            text=raw_value, highlight="yellow",
            note="Reviewer accepted this clause as originally requested — non-standard language; see review report.",
            source="slot",
        )]

    if action.action == ReviewerActionType.REMOVE:
        return [RenderedRun(text="", source="slot")]

    if action.action == ReviewerActionType.EDIT:
        edited = action.edited_text or ""
        return [RenderedRun(
            text=edited, highlight="yellow",
            note="Reviewer-edited text — required safeguards were not re-verified automatically.",
            source="slot",
        )]

    return [RenderedRun(text="", source="slot")]  # unreachable given ReviewerActionType's enum


def _substitute_section(
    section_text: str, placeholder_runs: dict[str, list[RenderedRun]]
) -> list[RenderedRun]:
    runs: list[RenderedRun] = []
    pos = 0
    for m in _PLACEHOLDER_TOKEN.finditer(section_text):
        literal_before = section_text[pos:m.start()]
        if literal_before:
            runs.append(RenderedRun(text=literal_before))
        name = m.group(1)
        runs.extend(placeholder_runs.get(name, [RenderedRun(text=m.group(0))]))
        pos = m.end()
    tail = section_text[pos:]
    if tail:
        runs.append(RenderedRun(text=tail))
    return runs


def _build_section_4_runs(section_text: str, special_runs: list[RenderedRun]) -> list[RenderedRun]:
    """
    Section 4 needs bespoke handling: its placeholder sits at the very end
    of the paragraph, after "...Section 5. ". When the special-clause
    content is empty, the trailing space before the (now-removed) token
    must go too, or the paragraph ends "Section 5. " with a stray space.
    """
    idx = section_text.index(SPECIAL_CLAUSE_TOKEN)
    before = section_text[:idx]
    after = section_text[idx + len(SPECIAL_CLAUSE_TOKEN):]

    is_empty = not any(r.text for r in special_runs)
    if is_empty:
        before = before.rstrip()

    runs = [RenderedRun(text=before)]
    if not is_empty:
        runs.extend(special_runs)
    if after:
        runs.append(RenderedRun(text=after))
    return runs


def _signature_block_paragraphs() -> list[RenderedParagraph]:
    lines = [
        "",
        "Disclosing Party:",
        "Name: _________________________",
        "Title: _________________________",
        "Date: _________________________",
        "",
        "Receiving Party:",
        "Name: _________________________",
        "Title: _________________________",
        "Date: _________________________",
    ]
    return [RenderedParagraph(runs=[RenderedRun(text=line)]) for line in lines]


# ---------------------------------------------------------------------------
# write_docx — mechanical rendering only, no decisions
# ---------------------------------------------------------------------------


def write_docx(rendered: RenderedDoc) -> bytes:
    import docx
    from docx.enum.text import WD_COLOR_INDEX
    from docx.shared import Pt, RGBColor

    doc = docx.Document()

    if rendered.banner:
        banner_para = doc.add_paragraph()
        banner_run = banner_para.add_run(rendered.banner)
        banner_run.bold = True
        banner_run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)

    title_para = doc.add_paragraph()
    title_run = title_para.add_run(rendered.title)
    title_run.bold = True
    title_run.font.size = Pt(14)

    for paragraph in rendered.paragraphs:
        p = doc.add_paragraph()
        for run_data in paragraph.runs:
            if not run_data.text:
                continue
            run = p.add_run(run_data.text)
            if run_data.highlight == "yellow":
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW
            elif run_data.highlight == "red":
                run.font.highlight_color = WD_COLOR_INDEX.RED
            if run_data.note:
                _attach_note(doc, p, run, run_data.note)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _attach_note(doc, paragraph, run, note: str) -> None:
    """Word comment (confirmed supported, spike S-6) with an inline-text
    fallback (D-32) if the comment API ever raises in a given environment —
    document generation must never fail because a note couldn't attach."""
    try:
        doc.add_comment(run, text=note, author="Zycus Contract Agent", initials="ZCA")
    except Exception:
        note_run = paragraph.add_run(f" [Reviewer note: {note}]")
        note_run.italic = True
