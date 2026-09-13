"""
QA Gate (architecture.md §3 #15, §9, D-33). The safety net: every check
below must pass before a .docx is ever handed back to a reviewer (P4's
POST /api/run only returns docx_base64 when QAReport.passed is True).

Nothing here *fixes* a problem — a failing check blocks the download and
surfaces in the trace. Silently patching the document would defeat the
point of having a gate at all.
"""
from __future__ import annotations

import io
import re

import docx as docx_lib

from core.models import (
    BusinessInputs,
    ClauseAssessment,
    FieldDecision,
    FieldStatus,
    QACheckResult,
    QAReport,
    Readiness,
    ReviewerAction,
)
from tools.assembler import RenderedDoc, _PLACEHOLDER_TOKEN, _resolve_special_clause_runs
from tools.template_store import Template

_Q1_PLACEHOLDER_PATTERN = re.compile(r"\[[A-Z][A-Z ,]+\]")
_Q2_FORBIDDEN_LITERALS = ["today's date", "n/a", "not applicable", "payment"]


def _normalize_ws(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def run_qa(
    docx_bytes: bytes,
    rendered: RenderedDoc,
    template: Template,
    inputs: BusinessInputs,
    decisions: list[FieldDecision],
    clause: ClauseAssessment | None,
    readiness: Readiness,
    reviewer_actions: list[ReviewerAction] | None = None,
) -> QAReport:
    decisions_by_field = {d.field: d for d in decisions}
    checks = [
        _check_q1_no_leftover_placeholders(rendered),
        _check_q2_no_forbidden_literals_in_slots(rendered),
        _check_q3_consistent_values(template, rendered, decisions_by_field),
        _check_q4_all_sections_present(template, rendered),
        _check_q5_non_slot_text_unchanged(template, rendered),
        _check_q6_clause_matches_decision(
            template, rendered, decisions_by_field.get("special_clause"), clause,
            reviewer_actions, inputs.special_clause,
        ),
        _check_q7_readiness_banner_consistent(rendered, readiness, decisions),
        _check_q8_docx_matches_rendered(docx_bytes, rendered),
    ]
    return QAReport(passed=all(c.passed for c in checks), checks=checks)


# --- Q1 ---------------------------------------------------------------------


def _check_q1_no_leftover_placeholders(rendered: RenderedDoc) -> QACheckResult:
    matches = _Q1_PLACEHOLDER_PATTERN.findall(rendered.full_text)
    return QACheckResult(
        check_id="Q1", name="No leftover template placeholders", passed=not matches,
        detail="clean" if not matches else f"found: {matches}",
    )


# --- Q2 ---------------------------------------------------------------------


def _check_q2_no_forbidden_literals_in_slots(rendered: RenderedDoc) -> QACheckResult:
    hits: list[str] = []
    for para in rendered.paragraphs:
        for run in para.runs:
            if run.source != "slot" or not run.text:
                continue
            lowered = run.text.lower()
            for literal in _Q2_FORBIDDEN_LITERALS:
                if literal in lowered:
                    hits.append(f"{literal!r} found in slot text {run.text!r}")
    return QACheckResult(
        check_id="Q2", name="No forbidden literals in slots", passed=not hits,
        detail="clean" if not hits else "; ".join(hits),
    )


# --- Q3 ---------------------------------------------------------------------


def _check_q3_consistent_values(
    template: Template, rendered: RenderedDoc, decisions_by_field: dict[str, FieldDecision]
) -> QACheckResult:
    mismatches: list[str] = []
    for placeholder_name, spec in template.placeholders.items():
        decision = decisions_by_field.get(spec.field)
        if decision is None or not decision.value_for_document:
            continue
        expected = decision.value_for_document
        for section in spec.sections:
            idx = template.sections.index(section)
            if idx >= len(rendered.paragraphs) or expected not in rendered.paragraphs[idx].text:
                mismatches.append(f"{placeholder_name} value inconsistent in section {section}")
    return QACheckResult(
        check_id="Q3", name="Values identical at every occurrence", passed=not mismatches,
        detail="clean" if not mismatches else "; ".join(mismatches),
    )


# --- Q4 ---------------------------------------------------------------------


def _check_q4_all_sections_present(template: Template, rendered: RenderedDoc) -> QACheckResult:
    n = len(template.sections)
    section_paragraphs = rendered.paragraphs[:n]
    problems = []
    if len(section_paragraphs) != n:
        problems.append(f"expected {n} section paragraphs, found {len(section_paragraphs)}")
    else:
        empty = [template.sections[i] for i, p in enumerate(section_paragraphs) if not p.text.strip()]
        if empty:
            problems.append(f"empty section(s): {empty}")
    return QACheckResult(
        check_id="Q4", name="All sections present, in order", passed=not problems,
        detail="clean" if not problems else "; ".join(problems),
    )


# --- Q5 ---------------------------------------------------------------------


def _check_q5_non_slot_text_unchanged(template: Template, rendered: RenderedDoc) -> QACheckResult:
    mismatches: list[str] = []
    n = len(template.sections)
    for i, section in enumerate(template.sections):
        if i >= len(rendered.paragraphs):
            mismatches.append(section)
            continue
        original = _PLACEHOLDER_TOKEN.sub("", template.section_text[section])
        original_norm = _normalize_ws(original)
        rendered_template_text = "".join(
            r.text for r in rendered.paragraphs[i].runs if r.source == "template"
        )
        rendered_norm = _normalize_ws(rendered_template_text)
        if original_norm != rendered_norm:
            mismatches.append(section)
    return QACheckResult(
        check_id="Q5", name="Non-slot text matches the template verbatim", passed=not mismatches,
        detail="clean" if not mismatches else f"sections reworded: {mismatches}",
    )


# --- Q6 ---------------------------------------------------------------------


def _check_q6_clause_matches_decision(
    template: Template,
    rendered: RenderedDoc,
    special_decision: FieldDecision | None,
    clause: ClauseAssessment | None,
    reviewer_actions: list[ReviewerAction] | None,
    raw_special_clause: str,
) -> QACheckResult:
    reviewer_action = next(
        (a for a in (reviewer_actions or []) if a.field == "special_clause"), None
    )
    expected_runs = _resolve_special_clause_runs(
        special_decision, clause, reviewer_action, raw_special_clause
    )
    expected_text = _normalize_ws("".join(r.text for r in expected_runs))

    section4_index = template.sections.index("4")
    actual_runs = [r for r in rendered.paragraphs[section4_index].runs if r.source == "slot"]
    actual_text = _normalize_ws("".join(r.text for r in actual_runs))

    passed = expected_text == actual_text
    return QACheckResult(
        check_id="Q6", name="§4 clause content matches the router/reviewer decision", passed=passed,
        detail="clean" if passed else f"expected {expected_text!r}, found {actual_text!r}",
    )


# --- Q7 ---------------------------------------------------------------------


def _check_q7_readiness_banner_consistent(
    rendered: RenderedDoc, readiness: Readiness, decisions: list[FieldDecision]
) -> QACheckResult:
    blocked_count = sum(1 for d in decisions if d.status == FieldStatus.BLOCKED_MISSING)
    marker_count = rendered.full_text.count("⟦MISSING:")

    problems = []
    if readiness == Readiness.BLOCKED:
        if not rendered.banner:
            problems.append("readiness is BLOCKED but no banner is present")
        if "NOT READY" not in (rendered.banner or ""):
            problems.append("banner does not say NOT READY")
        if marker_count != blocked_count:
            problems.append(f"{marker_count} MISSING markers but {blocked_count} blocked fields")
    else:
        if rendered.banner:
            problems.append(f"banner present but readiness is {readiness.value}")
        if marker_count or blocked_count:
            problems.append("MISSING markers/blocked fields present without BLOCKED readiness")

    return QACheckResult(
        check_id="Q7", name="Readiness banner consistent with blocked fields", passed=not problems,
        detail="clean" if not problems else "; ".join(problems),
    )


# --- Q8 ---------------------------------------------------------------------


def _check_q8_docx_matches_rendered(docx_bytes: bytes, rendered: RenderedDoc) -> QACheckResult:
    try:
        doc = docx_lib.Document(io.BytesIO(docx_bytes))
    except Exception as e:  # noqa: BLE001 — any failure to reopen is a Q8 failure, not a crash
        return QACheckResult(
            check_id="Q8", name=".docx reopens and matches RenderedDoc", passed=False,
            detail=f"failed to reopen .docx: {e!r}",
        )
    docx_text = _normalize_ws("\n".join(p.text for p in doc.paragraphs))
    rendered_text = _normalize_ws(rendered.full_text)
    passed = docx_text == rendered_text
    return QACheckResult(
        check_id="Q8", name=".docx reopens and matches RenderedDoc", passed=passed,
        detail="clean" if passed else "docx text does not match RenderedDoc.full_text after reopening",
    )
